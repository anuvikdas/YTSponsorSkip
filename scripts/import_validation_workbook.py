#!/usr/bin/env python3
import argparse
import csv
import hashlib
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

VIDEO_FIELDS = [
    "video_id",
    "url",
    "title",
    "caption_type",
    "review_status",
    "promotion_status",
    "split",
    "notes",
]
SEGMENT_FIELDS = [
    "video_id",
    "segment_number",
    "category",
    "start_seconds",
    "end_seconds",
    "boundary_certainty",
    "notes",
]

SPLIT_MAP = {"Development": "tune", "Held-out": "held_out"}
REVIEW_STATUS_MAP = {"Fully reviewed": "reviewed", "Partial": "in_progress"}
PROMOTION_STATUS_MAP = {"Present": "one_or_more", "None": "none", "Unknown": "unknown"}
CAPTION_TYPE_MAP = {
    "Manual": "manual",
    "Generated": "generated",
    "None": "none",
    "Unknown": "unknown",
}

EXPECTED_TOTALS = {
    "videos": 24,
    "promotion_videos": 12,
    "segments": 13,
    "negative_videos": 10,
    "partial_videos": 2,
}


def rows_by_header(worksheet: Any) -> list[dict[str, Any]]:
    values = list(worksheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows: list[dict[str, Any]] = []
    for values_row in values[1:]:
        row = dict(zip(headers, values_row, strict=True))
        if row.get("video_id") not in (None, ""):
            rows.append(row)
    return rows


def mapped(
    mapping: dict[str, str], value: Any, field: str, row_number: int, errors: list[str]
) -> str:
    if value in mapping:
        return mapping[value]
    errors.append(f"Videos row {row_number}: unrecognized {field} value {value!r}.")
    return ""


def finite_number(value: Any, field: str, row_number: int, errors: list[str]) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        errors.append(f"Segments row {row_number}: {field} must be numeric.")
        return None
    converted = float(value)
    if not math.isfinite(converted):
        errors.append(f"Segments row {row_number}: {field} must be finite.")
        return None
    return converted


def csv_number(value: float) -> int | float:
    return int(value) if value.is_integer() else value


def write_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as destination:
        writer = csv.DictWriter(destination, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def import_workbook(source: Path, output_dir: Path) -> dict[str, Any]:
    source_hash_before = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=False)
    errors: list[str] = []
    warnings: list[str] = []

    missing_sheets = {"Videos", "Segments"} - set(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Missing required sheets: {', '.join(sorted(missing_sheets))}")

    source_videos = rows_by_header(workbook["Videos"])
    source_segments = rows_by_header(workbook["Segments"])
    workbook.close()

    videos: list[dict[str, Any]] = []
    source_by_video: dict[str, dict[str, Any]] = {}
    seen_video_ids: set[str] = set()
    for row_number, row in enumerate(source_videos, start=2):
        video_id = str(row.get("video_id") or "").strip()
        if not video_id:
            errors.append(f"Videos row {row_number}: video_id is required.")
            continue
        if video_id in seen_video_ids:
            errors.append(f"Videos row {row_number}: duplicate video_id {video_id!r}.")
        seen_video_ids.add(video_id)
        source_by_video[video_id] = row
        videos.append(
            {
                "video_id": video_id,
                "url": row.get("url") or "",
                "title": row.get("title") or "",
                "caption_type": mapped(
                    CAPTION_TYPE_MAP, row.get("caption_type"), "caption_type", row_number, errors
                ),
                "review_status": mapped(
                    REVIEW_STATUS_MAP, row.get("review_status"), "review_status", row_number, errors
                ),
                "promotion_status": mapped(
                    PROMOTION_STATUS_MAP,
                    row.get("promotion_status"),
                    "promotion_status",
                    row_number,
                    errors,
                ),
                "split": mapped(SPLIT_MAP, row.get("split"), "split", row_number, errors),
                "notes": row.get("review_notes") or "",
            }
        )

    segments: list[dict[str, Any]] = []
    segment_keys: set[tuple[str, int]] = set()
    segments_by_video: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    duration_checks = 0
    for row_number, row in enumerate(source_segments, start=2):
        video_id = str(row.get("video_id") or "").strip()
        if video_id not in source_by_video:
            errors.append(f"Segments row {row_number}: unknown video_id {video_id!r}.")

        raw_segment_number = row.get("segment_number")
        if (
            isinstance(raw_segment_number, bool)
            or not isinstance(raw_segment_number, (int, float))
            or not float(raw_segment_number).is_integer()
            or raw_segment_number < 1
        ):
            errors.append(f"Segments row {row_number}: segment_number must be a positive integer.")
            continue
        segment_number = int(raw_segment_number)
        key = (video_id, segment_number)
        if key in segment_keys:
            errors.append(f"Segments row {row_number}: duplicate segment key {key!r}.")
        segment_keys.add(key)

        start = finite_number(row.get("start_seconds"), "start_seconds", row_number, errors)
        end = finite_number(row.get("end_seconds"), "end_seconds", row_number, errors)
        if start is None or end is None:
            continue
        if start < 0 or end < 0:
            errors.append(f"Segments row {row_number}: timestamps must be nonnegative.")
        if start >= end:
            errors.append(
                f"Segments row {row_number}: start_seconds must be less than end_seconds."
            )

        source_video = source_by_video.get(video_id, {})
        duration = source_video.get("duration_seconds")
        if duration not in (None, ""):
            duration_checks += 1
            numeric_duration = finite_number(duration, "video duration_seconds", row_number, errors)
            if numeric_duration is not None and end > numeric_duration:
                errors.append(
                    f"Segments row {row_number}: end_seconds exceeds video duration_seconds."
                )

        normalized = {
            "video_id": video_id,
            "segment_number": segment_number,
            "category": row.get("category") or "",
            "start_seconds": csv_number(start),
            "end_seconds": csv_number(end),
            "boundary_certainty": row.get("boundary_certainty") or "",
            "notes": row.get("notes") or "",
        }
        segments.append(normalized)
        segments_by_video[video_id].append(normalized)

    normalized_by_video = {row["video_id"]: row for row in videos}
    for video_id, video in normalized_by_video.items():
        video_segments = segments_by_video[video_id]
        if video["review_status"] == "in_progress":
            if video["promotion_status"] != "unknown":
                errors.append(
                    f"{video_id}: in-progress review must retain promotion_status=unknown."
                )
            if video_segments:
                errors.append(f"{video_id}: in-progress review unexpectedly has segment labels.")
        elif video["promotion_status"] == "none" and video_segments:
            errors.append(f"{video_id}: reviewed negative unexpectedly has segment labels.")
        elif video["promotion_status"] == "one_or_more" and not video_segments:
            errors.append(f"{video_id}: promotion_status=one_or_more has no segment labels.")

    actual_totals = {
        "videos": len(videos),
        "promotion_videos": sum(row["promotion_status"] == "one_or_more" for row in videos),
        "segments": len(segments),
        "negative_videos": sum(row["promotion_status"] == "none" for row in videos),
        "partial_videos": sum(row["review_status"] == "in_progress" for row in videos),
    }
    for label, expected in EXPECTED_TOTALS.items():
        actual = actual_totals[label]
        if actual != expected:
            errors.append(f"Expected {expected} {label}, found {actual}.")

    incomplete = {
        "unknown_caption_type": sum(row["caption_type"] == "unknown" for row in videos),
        "missing_review_date": sum(
            source_by_video[row["video_id"]].get("reviewed_on") in (None, "") for row in videos
        ),
        "missing_video_duration": sum(
            source_by_video[row["video_id"]].get("duration_seconds") in (None, "") for row in videos
        ),
        "missing_segment_category": sum(not row["category"] for row in segments),
        "missing_boundary_certainty": sum(not row["boundary_certainty"] for row in segments),
    }
    if duration_checks == 0:
        warnings.append(
            "No source video durations were supplied; timestamp bounds could not be checked "
            "against video duration."
        )

    source_hash_after = hashlib.sha256(source.read_bytes()).hexdigest()
    if source_hash_before != source_hash_after:
        errors.append("The source workbook changed during import.")

    write_csv(output_dir / "videos.csv", VIDEO_FIELDS, videos)
    write_csv(output_dir / "segments.csv", SEGMENT_FIELDS, segments)

    return {
        "source": str(source),
        "source_sha256": source_hash_after,
        "source_preserved": source_hash_before == source_hash_after,
        "actual_totals": actual_totals,
        "split_counts": dict(Counter(row["split"] for row in videos)),
        "incomplete_annotations": incomplete,
        "duration_checks_performed": duration_checks,
        "errors": errors,
        "warnings": warnings,
    }


def render_report(result: dict[str, Any]) -> str:
    totals = result["actual_totals"]
    incomplete = result["incomplete_annotations"]
    errors = result["errors"]
    warnings = result["warnings"]
    tune_count = result["split_counts"].get("tune", 0)
    held_out_count = result["split_counts"].get("held_out", 0)
    mapping_note = (
        "`Development` was mapped to `tune`; `Held-out` was mapped to `held_out`. "
        "`Fully reviewed`, `Partial`, `Present`, and `None` were mapped to the project "
        "enums without changing label meaning. Only `review_notes` was mapped into Videos "
        "`notes`; candidate-selection metadata was excluded from normalized labels."
    )
    incomplete_note = (
        "Blank category and boundary-certainty fields remain blank. Unknown caption types "
        "remain `unknown`. No review dates, durations, categories, or certainty values were "
        "invented. Held-out rows are retained for future evaluation but must not be used by "
        "development fixtures or detector tuning."
    )
    manual_note = (
        "These are user-supplied manual annotations. The import validates structure and "
        "consistency; it does not independently verify the labels against the videos."
    )
    error_lines = "\n".join(f"- {error}" for error in errors) or "- None."
    warning_lines = "\n".join(f"- {warning}" for warning in warnings) or "- None."
    return f"""# Validation-label import report

Source workbook: `{result["source"]}`  
Source SHA-256: `{result["source_sha256"]}`  
Source preserved unchanged: `{str(result["source_preserved"]).lower()}`

## Result

- Import errors: {len(errors)}
- Videos: {totals["videos"]}
- Videos with promotions: {totals["promotion_videos"]}
- Promotional intervals: {totals["segments"]}
- Reviewed negatives: {totals["negative_videos"]}
- Partial reviews retained as `in_progress` / `unknown`: {totals["partial_videos"]}
- Split counts: {tune_count} tune, {held_out_count} held_out
- Duration-bound checks performed: {result["duration_checks_performed"]}

{mapping_note}

{manual_note}

## Incomplete annotations, not import errors

- Unknown caption type: {incomplete["unknown_caption_type"]} videos
- Missing review date: {incomplete["missing_review_date"]} videos
- Missing video duration: {incomplete["missing_video_duration"]} videos
- Missing segment category: {incomplete["missing_segment_category"]} intervals
- Missing boundary certainty: {incomplete["missing_boundary_certainty"]} intervals

{incomplete_note}

## Validation errors

{error_lines}

## Warnings

{warning_lines}
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("data/validation"))
    parser.add_argument("--report", type=Path, default=Path("docs/validation-import-report.md"))
    args = parser.parse_args()

    result = import_workbook(args.source, args.output_dir)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(render_report(result), encoding="utf-8")
    print(render_report(result))
    if result["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
